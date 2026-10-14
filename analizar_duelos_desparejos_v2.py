"""Analizador experimental de duelos tácticos basado en coordenadas PitchAPI.

Uso:
    python analizar_duelos_desparejos_v2.py --fecha 12

Este archivo es deliberadamente independiente del motor de equipos, predictor y aprendizaje.
No reemplaza analizar_duelos_desparejos.py. Las posiciones se infieren de coordenadas de
alineaciones y se guardan en hojas de auditoría para validar la clasificación.
Los duelos son estimaciones por sector, no marcajes individuales confirmados.
"""
from __future__ import annotations

import argparse
import json
import re
import unicodedata
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

BASE = Path(__file__).resolve().parent
DATOS = BASE / "datos"
PARTIDOS_DIR = DATOS / "partidos"
HISTORICO = DATOS / "dataset_winning_pitchapi.csv"
LINEUPS_DIR = DATOS / "pitchapi" / "lineups"
SALIDA = DATOS / "duelos_desparejos_v2_fecha{fecha}.xlsx"
INICIO_CLAUSURA = pd.Timestamp("2026-07-23")

METRICAS_DEF = {
    "duels_won": ("Duelos ganados", 1.0),
    "interceptions": ("Intercepciones", 1.0),
    "tackles": ("Entradas / tackles", 0.8),
    "recoveries": ("Recuperaciones", 0.6),
    "clearances": ("Despejes", 0.5),
    "blocks": ("Bloqueos", 0.5),
}
METRICAS_DEL = {
    "take_ons_won": ("Regates exitosos", 1.0),
    "duels_won": ("Duelos ganados", 0.9),
    "progressive_carries": ("Conducciones progresivas", 0.8),
    "fouls_won": ("Faltas recibidas", 0.6),
    "chances_created": ("Ocasiones creadas", 0.6),
    "shots_on_target": ("Tiros al arco", 0.5),
}
MIN_MINUTOS_MUESTRA = 45
ULTIMAS_APARICIONES = 5


def normalizar(valor):
    texto = unicodedata.normalize("NFKD", str(valor or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", texto.lower()).strip()


def encontrar_evento(obj):
    if isinstance(obj, dict):
        e = obj.get("event")
        if isinstance(e, dict) and "homeTeam" in e and "awayTeam" in e:
            return e
        for v in obj.values():
            r = encontrar_evento(v)
            if r:
                return r
    elif isinstance(obj, list):
        for v in obj:
            r = encontrar_evento(v)
            if r:
                return r
    return None


def partidos_objetivo(fecha):
    filas = []
    for path in PARTIDOS_DIR.glob("*.json"):
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            continue
        e = encontrar_evento(payload)
        if not e:
            continue
        round_no = (e.get("roundInfo") or {}).get("round")
        ts = e.get("startTimestamp")
        dt = pd.to_datetime(ts, unit="s", errors="coerce", utc=True) if ts else pd.NaT
        if pd.isna(dt) or dt.tz_convert(None) < INICIO_CLAUSURA or str(round_no) != str(fecha):
            continue
        home = (e.get("homeTeam") or {}).get("name")
        away = (e.get("awayTeam") or {}).get("name")
        if home and away:
            filas.append({"match_id": str(e.get("id") or path.stem),
                          "fecha_partido": dt.tz_convert(None),
                          "local": str(home), "visitante": str(away)})
    if not filas:
        raise ValueError(f"No encontré partidos de la Fecha {fecha} del Clausura en datos/partidos/*.json.")
    return pd.DataFrame(filas).drop_duplicates("match_id").sort_values("fecha_partido")


def cargar_json(path):
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return None
    data = raw.get("data", raw)
    return data if isinstance(data, dict) else None


def cargar_lineup(match_id):
    data = cargar_json(LINEUPS_DIR / f"{match_id}_lineups.json")
    if not data:
        return None
    if data.get("confirmed") is False:
        home_ok = bool((data.get("home") or {}).get("confirmed"))
        away_ok = bool((data.get("away") or {}).get("confirmed"))
        if not (home_ok and away_ok):
            return None
    return data


def lineup_side_name(lineup, side):
    value = lineup.get(f"{side}_team")
    if isinstance(value, dict):
        return str(value.get("name") or "")
    return str(value or "")


def side_for_team(lineup, club):
    target = normalizar(club)
    for side in ("home", "away"):
        if normalizar(lineup_side_name(lineup, side)) == target:
            return side
    return None


def lineup_players(lineup, club):
    side = side_for_team(lineup, club)
    if not side:
        return None
    team_data = lineup.get(side) or {}
    starters = team_data.get("starters")
    if not isinstance(starters, list) or not starters:
        return None
    result = {}
    for p in starters:
        if not isinstance(p, dict):
            continue
        pid = p.get("player_id", p.get("id"))
        if pid is None:
            continue
        try:
            x, y = float(p.get("pitch_x")), float(p.get("pitch_y"))
        except (TypeError, ValueError):
            x = y = np.nan
        result[str(pid)] = {
            "id": str(pid),
            "name": str(p.get("name") or p.get("player_name") or ""),
            "position_id": p.get("position_id"),
            "x": x, "y": y,
        }
    return result


def sofa_posiciones_por_partido():
    """Mapa match_id -> equipo -> nombre normalizado -> posición general SofaScore."""
    resultado = {}
    for path in PARTIDOS_DIR.glob("*.json"):
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            continue
        mid = str(path.stem)
        lineups = payload.get("lineups")
        if not isinstance(lineups, dict):
            continue
        por_equipo = {}
        for side in ("home", "away"):
            team_data = lineups.get(side)
            if not isinstance(team_data, dict):
                continue
            # SofaScore puede devolver el nombre del equipo en varios campos.
            team_obj = team_data.get("team") or team_data.get("name") or {}
            team_name = team_obj.get("name", "") if isinstance(team_obj, dict) else str(team_obj)
            mapping = {}
            for item in team_data.get("players", []) or []:
                player = item.get("player", {}) if isinstance(item, dict) else {}
                name = normalizar(player.get("name"))
                pos = str(player.get("position") or "").upper()
                if name and pos:
                    mapping[name] = pos
            if team_name:
                por_equipo[normalizar(team_name)] = mapping
            por_equipo[f"__{side}__"] = mapping
        if por_equipo:
            resultado[mid] = por_equipo
    return resultado


def calibrar_eje_largo(cache_lineups, sofa):
    """El eje longitudinal es el que separa más al arquero de la mediana de sus compañeros."""
    separaciones = {"x": [], "y": []}
    muestras = 0
    for mid, lineup in cache_lineups.items():
        if not lineup:
            continue
        for side in ("home", "away"):
            team_data = lineup.get(side) or {}
            team_name = normalizar(lineup_side_name(lineup, side))
            sofa_map = sofa.get(str(mid), {}).get(f"__{side}__", {})
            starters = team_data.get("starters") or []
            puntos = []
            arquero = None
            for p in starters:
                if not isinstance(p, dict):
                    continue
                name = normalizar(p.get("name") or p.get("player_name"))
                try:
                    x, y = float(p.get("pitch_x")), float(p.get("pitch_y"))
                except (TypeError, ValueError):
                    continue
                # Coordenadas pueden venir normalizadas 0..1 o como porcentaje 0..100.
                if max(abs(x), abs(y)) > 1.5:
                    x, y = x / 100.0, y / 100.0
                puntos.append((x, y))
                if str(p.get("position_id", "")) == "11" or sofa_map.get(name) == "G":
                    arquero = (x, y)
            if arquero and len(puntos) >= 6:
                otros = [p for p in puntos if p != arquero]
                separaciones["x"].append(abs(arquero[0] - float(np.median([p[0] for p in otros]))))
                separaciones["y"].append(abs(arquero[1] - float(np.median([p[1] for p in otros]))))
                muestras += 1
    med = {axis: float(np.median(vals)) if vals else 0.0 for axis, vals in separaciones.items()}
    if muestras < 10 or max(med.values()) < 0.12:
        return None, med, muestras
    return max(med, key=med.get), med, muestras


def preparar_historico(df):
    alias = {
        "player_id": ["player_id"],
        "nombre": ["player_name", "nombre_jugador", "name"],
        "equipo": ["team_name", "club", "team", "equipo"],
        "partido": ["match_id", "fixture_id", "partido_id"],
        "fecha": ["match_date", "date", "fecha", "fecha_partido"],
        "minutos": ["minutes_played", "minutes"],
    }
    cols = {dest: next((c for c in opts if c in df.columns), None) for dest, opts in alias.items()}
    missing = [k for k, v in cols.items() if v is None]
    if missing:
        raise ValueError("Al histórico le faltan columnas: " + ", ".join(missing))
    df = df.copy()
    for dest, original in cols.items():
        if dest != original:
            df[dest] = df[original]
    df["player_id"] = df["player_id"].astype(str)
    df["partido"] = df["partido"].astype(str)
    df["equipo"] = df["equipo"].astype(str)
    df["fecha"] = pd.to_datetime(df["fecha"], errors="coerce", utc=True).dt.tz_convert(None)
    df["minutos"] = pd.to_numeric(df["minutos"], errors="coerce").fillna(0)
    for col in set(METRICAS_DEF) | set(METRICAS_DEL):
        if col not in df:
            df[col] = np.nan
        df[col] = pd.to_numeric(df[col], errors="coerce")
    return df


def partidos_previos(df, team, fecha_corte):
    bloque = df[(df["equipo"].map(normalizar) == normalizar(team)) & (df["fecha"] < fecha_corte)]
    return bloque[["partido", "fecha"]].dropna().drop_duplicates().sort_values(
        ["fecha", "partido"], ascending=False).head(2)


def validar_titularidad_2de2(df, player, team, fecha_corte, cache):
    """Exige titularidad comprobada en los dos últimos partidos del equipo.

    El ID de jugador puede diferir entre el CSV estadístico y el lineup. Por eso
    se permite validar por nombre normalizado solo si coincide exactamente con
    un único titular de ese equipo en ese partido. Nunca se infiere titularidad
    a partir de minutos jugados o de la mera presencia en el histórico.
    """
    previos = partidos_previos(df, team, fecha_corte)
    if len(previos) != 2:
        return False, [], "No hay exactamente dos partidos previos identificables del equipo"
    datos_jugador = []
    for _, row in previos.iterrows():
        mid = str(row["partido"])
        if mid not in cache:
            cache[mid] = cargar_lineup(mid)
        lineup = cache[mid]
        if not lineup:
            return False, [], f"Falta lineup confirmado del partido previo {mid}"
        players = lineup_players(lineup, team)
        if players is None:
            return False, [], f"No se pudo verificar equipo/titulares en {mid}"

        titular = players.get(str(player))
        if titular is None:
            # Algunos archivos históricos usan IDs distintos para el mismo jugador.
            # Solo aceptar coincidencia exacta de nombre contra la lista de TITULARES.
            filas_jugador = df[
                (df["player_id"].astype(str) == str(player)) &
                (df["partido"].astype(str) == mid) &
                (df["equipo"].map(normalizar) == normalizar(team))
            ]
            nombres = {
                normalizar(n) for n in filas_jugador["nombre"].dropna().astype(str)
                if normalizar(n)
            }
            coincidencias = [
                p for p in players.values()
                if normalizar(p.get("name", "")) in nombres
            ]
            if len(coincidencias) == 1:
                titular = coincidencias[0]
            else:
                return False, [], (
                    f"No se confirmó como titular en el partido previo {mid} "
                    "(ID/nombre no coincide de forma única con los titulares)"
                )
        datos_jugador.append(titular)

    if len(datos_jugador) != 2:
        return False, datos_jugador, "No se pudo confirmar titularidad en los dos partidos previos"
    if any(not np.isfinite(p["x"]) or not np.isfinite(p["y"]) for p in datos_jugador):
        return False, datos_jugador, "Titular 2/2 pero sin coordenadas válidas en ambos lineups"
    return True, datos_jugador, "Titular confirmado 2/2 con coordenadas"


def coordenadas_normalizadas(item):
    x, y = item.get("x"), item.get("y")
    if not np.isfinite(x) or not np.isfinite(y):
        return None
    if max(abs(x), abs(y)) > 1.5:
        x, y = x / 100.0, y / 100.0
    if not (0 <= x <= 1 and 0 <= y <= 1):
        return None
    return x, y


def rol_por_coordenadas(lineup, club, player, eje_largo, sofa_map):
    """Infiere central/lateral/extremo usando el arquero identificado en SofaScore."""
    side = side_for_team(lineup, club)
    if not side or eje_largo not in ("x", "y"):
        return None
    equipo = lineup.get(side) or {}
    starters = equipo.get("starters") or []
    puntos = []
    arquero = None
    objetivo = None
    for p in starters:
        if not isinstance(p, dict):
            continue
        pid = p.get("player_id", p.get("id"))
        coords = coordenadas_normalizadas({"x": p.get("pitch_x"), "y": p.get("pitch_y")})
        if not coords:
            continue
        x, y = coords
        longitud, banda = (x, y) if eje_largo == "x" else (y, x)
        puntos.append((longitud, banda))
        nombre = normalizar(p.get("name") or p.get("player_name"))
        if str(p.get("position_id", "")) == "11" or sofa_map.get(nombre) == "G":
            arquero = (longitud, banda)
        if str(pid) == str(player):
            objetivo = (longitud, banda, str(p.get("name") or p.get("player_name") or ""))
    if not objetivo or len(puntos) < 7:
        return None
    # La orientación se obtiene del arquero identificado por posición G en SofaScore; no se adivina el extremo.
    if arquero is None:
        return None
    # Orientar profundidad desde la propia portería hacia el campo rival.
    extremos = [p[0] for p in puntos]
    mediana_banda = float(np.median([p[1] for p in puntos]))
    span = max(abs(max(extremos) - min(extremos)), 0.25)
    profundidad = (objetivo[0] - arquero[0]) / span
    # Si el arquero está en el extremo alto, invertir la profundidad para que propia área sea 0.
    if arquero[0] > float(np.median(extremos)):
        profundidad = -profundidad
    profundidad = max(0.0, min(1.0, profundidad))
    ancho = objetivo[1]
    distancia_centro = abs(ancho - mediana_banda)
    # Central: eje central. Banda: lejos del centro. Umbrales conservadores.
    es_central = distancia_centro <= 0.16
    if profundidad <= 0.43:
        if es_central:
            return {"sector": "CENTRAL_DEF", "rol": "Defensor central", "profundidad": profundidad, "banda": ancho}
        return {"sector": "LATERAL", "rol": "Lateral (lado de cancha)", "profundidad": profundidad, "banda": ancho}
    if profundidad >= 0.62:
        if es_central:
            return {"sector": "CENTRAL_DEL", "rol": "Delantero central", "profundidad": profundidad, "banda": ancho}
        return {"sector": "EXTREMO", "rol": "Extremo (lado de cancha)", "profundidad": profundidad, "banda": ancho}
    return {"sector": "INTERMEDIO", "rol": "Zona intermedia; no clasifica como DEF/DEL", "profundidad": profundidad, "banda": ancho}


def metricas_por90(df, pid, team, fecha_corte):
    bloque = df[(df["player_id"] == str(pid)) &
                (df["equipo"].map(normalizar) == normalizar(team)) &
                (df["fecha"] < fecha_corte)]
    bloque = bloque.sort_values(["fecha", "partido"]).drop_duplicates("partido", keep="last").tail(ULTIMAS_APARICIONES)
    minutos = float(bloque["minutos"].sum())
    if minutos < MIN_MINUTOS_MUESTRA:
        return None
    out = {"minutos_muestra": minutos, "partidos_muestra": int(bloque["partido"].nunique())}
    for col in set(METRICAS_DEF) | set(METRICAS_DEL):
        total = bloque[col].sum(min_count=1)
        out[col] = float(total) / minutos * 90 if pd.notna(total) else np.nan
    names = bloque["nombre"].dropna().astype(str)
    out["nombre"] = names.iloc[-1] if len(names) else str(pid)
    return out


def vector(datos, metricas):
    vals = [(max(0.0, float(datos[c])), peso) for c, (_, peso) in metricas.items()
            if c in datos and pd.notna(datos[c])]
    return sum(np.log1p(v) * w for v, w in vals) / sum(w for _, w in vals) if vals else np.nan


def detalle(datos, metricas):
    return "; ".join(f"{nombre}: {datos[c]:.2f}/90" for c, (nombre, _) in metricas.items()
                     if c in datos and pd.notna(datos[c]))


def justificar_duelo(defensor, atacante, diferencia):
    """Explica en lenguaje sencillo qué inclina la comparación estadística."""

    def estadisticas_destacadas(datos, metricas, limite=2):
        valores = []
        for clave, (nombre, _) in metricas.items():
            valor = datos.get(clave)
            if valor is not None and pd.notna(valor):
                valores.append((nombre.lower(), float(valor)))
        valores.sort(key=lambda par: par[1], reverse=True)
        return ", ".join(
            f"{nombre} ({valor:.2f} por 90 min)"
            for nombre, valor in valores[:limite]
        )

    nombre_def = defensor.get("nombre", "El defensor")
    nombre_del = atacante.get("nombre", "El atacante")
    brecha = abs(float(diferencia))

    if diferencia > 0:
        inicio = (
            f"{nombre_def} sale favorecido en esta comparación frente a {nombre_del}: "
            f"su índice defensivo supera al índice ofensivo del atacante por "
            f"{brecha:.3f} puntos."
        )
        lectura = (
            "Dicho en criollo, los números recientes inclinan el cruce del lado "
            "del defensor: dentro de este modelo, sus estadísticas defensivas "
            "pesan más que las ofensivas del rival."
        )
    else:
        inicio = (
            f"{nombre_del} sale favorecido en esta comparación frente a {nombre_def}: "
            f"su índice ofensivo supera al índice defensivo del defensor por "
            f"{brecha:.3f} puntos."
        )
        lectura = (
            "Dicho en criollo, los números recientes inclinan el cruce del lado "
            "del atacante: dentro de este modelo, sus estadísticas ofensivas "
            f"pesan más que las defensivas de {nombre_def}."
        )

    detalles = []
    def_top = estadisticas_destacadas(defensor, METRICAS_DEF)
    del_top = estadisticas_destacadas(atacante, METRICAS_DEL)
    if def_top:
        detalles.append(f"En defensa, {nombre_def} destaca en {def_top}.")
    if del_top:
        detalles.append(f"En ataque, {nombre_del} destaca en {del_top}.")

    cierre = (
        "El cruce se seleccionó porque ambos jugadores encajan en sectores "
        "compatibles según las coordenadas y la diferencia entre sus índices "
        "supera el umbral del análisis. Es una estimación basada en partidos "
        "anteriores: no confirma un marcaje individual ni garantiza lo que "
        "pasará durante el partido."
    )
    return " ".join([inicio, lectura] + detalles + [cierre])


def banda_fisica(valor, centro):
    return "BANDA_A" if valor < centro else "BANDA_B"


def main():
    parser = argparse.ArgumentParser(description="Analiza duelos estimados con coordenadas PitchAPI; no modifica el motor.")
    parser.add_argument("--fecha", type=int, required=True)
    args = parser.parse_args()

    if not HISTORICO.exists():
        raise FileNotFoundError(f"No existe el histórico: {HISTORICO}")
    historico = preparar_historico(pd.read_csv(HISTORICO, low_memory=False))
    fixtures = partidos_objetivo(args.fecha)
    cache = {}
    # Cargar lineups históricos disponibles para calibrar el eje y la orientación.
    for p in LINEUPS_DIR.glob("*_lineups.json"):
        mid = p.name.replace("_lineups.json", "")
        cache[mid] = cargar_lineup(mid)
    sofa = sofa_posiciones_por_partido()
    eje_largo, separaciones, muestras_calibracion = calibrar_eje_largo(cache, sofa)
    if eje_largo is None:
        raise RuntimeError(
            "No se pudo calibrar de forma fiable el eje longitudinal con arqueros de SofaScore y coordenadas PitchAPI. "
            "No se generó un ranking para evitar inventar sectores. Revisá los lineups/posiciones."
        )

    salidas, descartes, auditoria = [], [], []
    for _, fx in fixtures.iterrows():
        home, away, fecha_corte = fx["local"], fx["visitante"], fx["fecha_partido"]
        for team_attack, team_defend in ((home, away), (away, home)):
            del_list, def_list = [], []
            for team, grupo, lista in ((team_attack, "DEL", del_list), (team_defend, "DEF", def_list)):
                rows = historico[(historico["equipo"].map(normalizar) == normalizar(team)) &
                                 (historico["fecha"] < fecha_corte)]
                pids = rows["player_id"].dropna().astype(str).unique()
                for pid in pids:
                    ok, previous, reason = validar_titularidad_2de2(historico, pid, team, fecha_corte, cache)
                    if not ok:
                        # Registrar motivo de exclusión para auditoría; no silenciar descartes.
                        name_rows = rows[rows["player_id"] == pid]["nombre"].dropna().astype(str)
                        descartes.append({"Partido": f"{home} vs {away}", "Equipo": team,
                                          "Jugador": name_rows.iloc[-1] if len(name_rows) else pid, "Motivo": reason})
                        continue
                    # Para el rol táctico usamos la alineación más reciente previa, no el partido futuro.
                    roles_previos = []
                    for prev_row in partidos_previos(historico, team, fecha_corte).itertuples(index=False):
                        prev_mid_role = str(prev_row.partido)
                        if prev_mid_role not in cache:
                            cache[prev_mid_role] = cargar_lineup(prev_mid_role)
                        prev_lineup = cache.get(prev_mid_role)
                        prev_side = side_for_team(prev_lineup, team) if prev_lineup else None
                        sofa_map = sofa.get(prev_mid_role, {}).get(f"__{prev_side}__", {}) if prev_side else {}
                        roles_previos.append(rol_por_coordenadas(cache.get(prev_mid_role), team, pid, eje_largo, sofa_map))
                    role = roles_previos[0] if roles_previos else None
                    if len(roles_previos) != 2 or any(r is None for r in roles_previos):
                        descartes.append({"Partido": f"{home} vs {away}", "Equipo": team,
                                          "Jugador": pid, "Motivo": "No se pudo inferir el sector en ambos lineups previos"})
                        continue
                    if roles_previos[0]["sector"] != roles_previos[1]["sector"]:
                        descartes.append({"Partido": f"{home} vs {away}", "Equipo": team,
                                          "Jugador": pid, "Motivo": "Cambió de sector entre los dos lineups previos"})
                        continue
                    stats = metricas_por90(historico, pid, team, fecha_corte)
                    if stats is None:
                        descartes.append({"Partido": f"{home} vs {away}", "Equipo": team,
                                          "Jugador": pid, "Motivo": "Menos de 45 minutos en la ventana estadística previa"})
                        continue
                    if role is None:
                        descartes.append({"Partido": f"{home} vs {away}", "Equipo": team,
                                          "Jugador": stats["nombre"], "Motivo": "No se pudo inferir rol con coordenadas válidas"})
                        continue
                    audit_row = {"Partido": f"{home} vs {away}", "Equipo": team, "Jugador": stats["nombre"],
                                 "Player ID": pid, "Grupo evaluado": grupo, "Eje longitudinal": eje_largo,
                                 "Rol inferido": role["rol"], "Sector": role["sector"],
                                 "Profundidad normalizada": round(role["profundidad"], 3),
                                 "Coordenada lateral": round(role["banda"], 3),
                                 "Titularidad 2/2": "Sí"}
                    auditoria.append(audit_row)
                    if grupo == "DEL":
                        if role["sector"] in ("CENTRAL_DEL", "EXTREMO"):
                            lista.append((pid, stats, role, audit_row))
                        else:
                            descartes.append({"Partido": f"{home} vs {away}", "Equipo": team,
                                              "Jugador": stats["nombre"], "Motivo": "Titular pero coordenadas no lo ubican en sector delantero"})
                    else:
                        if role["sector"] in ("CENTRAL_DEF", "LATERAL"):
                            lista.append((pid, stats, role, audit_row))
                        else:
                            descartes.append({"Partido": f"{home} vs {away}", "Equipo": team,
                                              "Jugador": stats["nombre"], "Motivo": "Titular pero coordenadas no lo ubican en sector defensivo"})

            # Match center-to-center; wide players match opposing fullbacks in the same physical lane.
            for pid_del, ds, rd, ad in del_list:
                for pid_def, fs, rf, af in def_list:
                    compatible = None
                    if rd["sector"] == "CENTRAL_DEL" and rf["sector"] == "CENTRAL_DEF":
                        compatible = "Delantero central vs defensor central"
                    elif rd["sector"] == "EXTREMO" and rf["sector"] == "LATERAL":
                        # Compare lateral coordinates; sides must be on the same physical flank.
                        center_del = 0.5
                        center_def = 0.5
                        side_del = banda_fisica(rd["banda"], center_del)
                        side_def = banda_fisica(rf["banda"], center_def)
                        if side_del == side_def:
                            compatible = f"Extremo vs lateral ({side_del})"
                    if not compatible:
                        continue
                    fdef, fdel = vector(fs, METRICAS_DEF), vector(ds, METRICAS_DEL)
                    if pd.isna(fdef) or pd.isna(fdel):
                        continue
                    diff = fdef - fdel
                    if abs(diff) < 0.12:
                        continue
                    favorecido = fs["nombre"] if diff > 0 else ds["nombre"]
                    salidas.append({
                        "Fecha": args.fecha, "Partido": f"{home} vs {away}",
                        "Fecha partido": fecha_corte.strftime("%d/%m/%Y"),
                        "Equipo defensor": team_defend, "Defensor": fs["nombre"],
                        "Rol defensor": rf["rol"], "Equipo atacante": team_attack,
                        "Atacante": ds["nombre"], "Rol atacante": rd["rol"],
                        "Sector compatible": compatible,
                        "Ventaja": "Defensor con ventaja estadística" if diff > 0 else "Atacante con ventaja estadística",
                        "Jugador favorecido": favorecido, "Índice defensor": round(fdef, 3),
                        "Índice atacante": round(fdel, 3), "Diferencia índice": round(diff, 3),
                        "Muestra defensor": f'{fs["partidos_muestra"]} partidos / {fs["minutos_muestra"]:.0f} min',
                        "Muestra atacante": f'{ds["partidos_muestra"]} partidos / {ds["minutos_muestra"]:.0f} min',
                        "Stats defensor por 90": detalle(fs, METRICAS_DEF),
                        "Stats atacante por 90": detalle(ds, METRICAS_DEL),
                        "Advertencia": "DUELO ESTIMADO por coordenadas; no confirma marcaje individual",
                        "Justificacion": justificar_duelo(fs, ds, diff),
                    })

    columnas = [
        "Fecha", "Partido", "Equipo defensor", "Defensor", "Equipo atacante",
        "Atacante", "Jugador favorecido", "Diferencia índice", "Justificacion"
    ]
    resultado = pd.DataFrame(salidas, columns=columnas)
    if not resultado.empty:
        resultado["_abs"] = resultado["Diferencia índice"].abs()
        resultado = resultado.sort_values(["_abs", "Partido"], ascending=[False, True]).drop(columns="_abs")
    resumen = pd.DataFrame([{
        "Fecha analizada": args.fecha, "Partidos encontrados": len(fixtures),
        "Duelos compatibles encontrados": len(resultado), "Eje longitudinal inferido": eje_largo,
        "Separación mediana arquero-compañeros en X": round(separaciones["x"], 3),
        "Separación mediana arquero-compañeros en Y": round(separaciones["y"], 3),
        "Muestras usadas para calibrar": muestras_calibracion,
        "Regla titularidad": "Titular confirmado en ambos partidos previos a la fecha objetivo",
        "Ventana estadística": f"Últimas {ULTIMAS_APARICIONES} apariciones anteriores al partido; mínimo {MIN_MINUTOS_MUESTRA} minutos",
        "Nota": "Roles inferidos por coordenadas; revisar hoja Auditoria_roles antes de interpretar resultados."
    }])
    salida = SALIDA.with_name(SALIDA.name.format(fecha=args.fecha))
    salida.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(salida, engine="openpyxl") as writer:
        resumen.to_excel(writer, sheet_name="Resumen", index=False)
        resultado.to_excel(writer, sheet_name="Duelos", index=False)
        pd.DataFrame(descartes, columns=["Partido", "Equipo", "Jugador", "Motivo"]).to_excel(writer, sheet_name="Exclusiones", index=False)
        pd.DataFrame(auditoria).to_excel(writer, sheet_name="Auditoria_roles", index=False)
    wb = load_workbook(salida)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="001E5F")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 32
        for col in ws.columns:
            encabezado = str(col[0].value or "")
            max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            if encabezado == "Justificacion":
                ws.column_dimensions[col[0].column_letter].width = 75
                for fila in range(2, ws.max_row + 1):
                    celda = ws.cell(fila, col[0].column)
                    celda.alignment = Alignment(vertical="top", wrap_text=True)
                    texto_celda = str(celda.value or "")
                    lineas = max(1, (len(texto_celda) // 90) + 1)
                    ws.row_dimensions[fila].height = min(max(lineas * 30, 60), 150)
            else:
                ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 12), 55)
    wb.save(salida)
    print(f"Fecha {args.fecha}: {len(fixtures)} partidos, {len(resultado)} duelos compatibles")
    print(f"Eje longitudinal inferido: {eje_largo}; separaciones medianas: {separaciones}; muestras: {muestras_calibracion}")
    print(f"Excel: {salida}")
    print("Importante: revisar Auditoria_roles antes de usar los duelos; no se modificó el motor de equipos.")


if __name__ == "__main__":
    main()
