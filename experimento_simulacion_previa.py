import sys
import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, Border, Side, PatternFill
import backtest_fecha10 as motor

FECHA_OBJETIVO = motor.FECHA_OBJETIVO
CORTE = motor.CORTE_HISTORICO
N_SIM = 10000
SEED = 42

SALIDA_CANDIDATOS = f"datos/fecha{FECHA_OBJETIVO}_pre_simulacion_candidatos.csv"
SALIDA_EQUIPOS = f"datos/fecha{FECHA_OBJETIVO}_pre_simulacion_equipos.csv"
SALIDA_FLEX = f"datos/fecha{FECHA_OBJETIVO}_pre_simulacion_flex.csv"
SALIDA_EXCEL = f"datos/fecha{FECHA_OBJETIVO}_equipos_predichos_excel.xlsx"


def simular_distribucion(valores, rng):
    valores = pd.to_numeric(pd.Series(valores), errors="coerce").dropna()
    if valores.empty:
        return {
            "pre_sim_n": 0, "pre_sim_media": np.nan, "pre_sim_p50": np.nan,
            "pre_sim_p75": np.nan, "pre_sim_p90": np.nan, "pre_sim_p95": np.nan,
            "pre_sim_std": np.nan, "pre_sim_prob_10": np.nan, "pre_sim_prob_15": np.nan,
        }
    muestra = rng.choice(valores.to_numpy(float), size=N_SIM, replace=True)
    return {
        "pre_sim_n": int(len(valores)),
        "pre_sim_media": float(np.mean(muestra)),
        "pre_sim_p50": float(np.quantile(muestra, .50)),
        "pre_sim_p75": float(np.quantile(muestra, .75)),
        "pre_sim_p90": float(np.quantile(muestra, .90)),
        "pre_sim_p95": float(np.quantile(muestra, .95)),
        "pre_sim_std": float(np.std(muestra)),
        "pre_sim_prob_10": float(np.mean(muestra >= 10)),
        "pre_sim_prob_15": float(np.mean(muestra >= 15)),
    }


def agregar_pre_simulacion(candidatos, historico):
    rng = np.random.default_rng(SEED)
    h = historico.copy()
    h["date"] = pd.to_datetime(h["date"], errors="coerce")
    h["winning_total"] = pd.to_numeric(h["winning_total"], errors="coerce")
    h = h[h["date"].notna() & (h["date"] < CORTE) & h["winning_total"].notna()].copy()

    filas = []
    for _, j in candidatos.iterrows():
        pid = str(j["player_id"])
        club = str(j["team_name"])
        serie = h[(h["player_id"].astype(str) == pid) & (h["team_name"].astype(str) == club)]["winning_total"]
        if serie.empty:
            serie = h[h["player_id"].astype(str) == pid]["winning_total"]
        filas.append(simular_distribucion(serie, rng))

    m = pd.DataFrame(filas, index=candidatos.index)
    for c in m.columns:
        candidatos[c] = m[c]

    media = pd.to_numeric(candidatos["pre_sim_media"], errors="coerce")
    contexto = pd.to_numeric(candidatos["score_contextual"], errors="coerce")
    matchup = pd.to_numeric(candidatos["matchup_score"], errors="coerce").clip(0.0, 1.0)

    # Influencia continua: 0.50 = neutro; 0.00 = -30%; 1.00 = +30%.
    factor_matchup = 1.0 + (matchup.fillna(0.50) - 0.50) * 0.60
    centro_contextual = (contexto * factor_matchup).fillna(media)
    delta = (centro_contextual - media).fillna(0.0)

    for c in ["pre_sim_media", "pre_sim_p50", "pre_sim_p75", "pre_sim_p90", "pre_sim_p95"]:
        candidatos[c + "_ajustada"] = pd.to_numeric(candidatos[c], errors="coerce") + delta

    candidatos["pre_sim_factor_matchup"] = factor_matchup
    candidatos["pre_sim_centro_contextual"] = centro_contextual
    candidatos["score_pre_sim_SEGURO"] = candidatos["pre_sim_p50_ajustada"]
    candidatos["score_pre_sim_INTERMEDIO"] = candidatos["pre_sim_p75_ajustada"]
    candidatos["score_pre_sim_ARRIESGADO"] = candidatos["pre_sim_p90_ajustada"]
    return candidatos


def activar_score_pre_simulacion():
    original = motor.calcular_score_seleccion

    def score_experimental(df, perfil):
        r = original(df, perfil).copy()
        columna = "score_pre_sim_" + perfil
        r["score_seleccion"] = pd.to_numeric(r[columna], errors="coerce").fillna(-np.inf)
        r["score_seleccion_original"] = r["score_seleccion"]
        return r

    motor.calcular_score_seleccion = score_experimental


def guardar_equipos(equipos):
    filas = []
    for perfil, jugadores in equipos.items():
        for j in jugadores:
            fila = dict(j)
            fila["motor"] = "PRE_SIMULACION"
            fila["perfil"] = perfil
            fila["tipo_registro"] = "TITULAR"
            filas.append(fila)
    pd.DataFrame(filas).to_csv(SALIDA_EQUIPOS, index=False, encoding="utf-8-sig")


def construir_flex(candidatos, equipos):
    flex_por_perfil = {}
    flex_usados = set()
    titulares_global = {str(j["player_id"]) for js in equipos.values() for j in js}

    for perfil in ["SEGURO", "INTERMEDIO", "ARRIESGADO"]:
        propios = {str(j["player_id"]) for j in equipos.get(perfil, [])}
        otros = titulares_global - propios
        flex = motor.construir_flex(
            candidatos.copy(), perfil, propios, otros,
            jugadores_titulares_global=titulares_global,
            flex_usados_global=flex_usados,
        )
        flex_por_perfil[perfil] = flex
        flex_usados.update(str(j["player_id"]) for j in flex)

    filas = []
    for perfil, jugadores in flex_por_perfil.items():
        for j in jugadores:
            fila = dict(j)
            fila["motor"] = "PRE_SIMULACION"
            fila["perfil"] = perfil
            fila["tipo_registro"] = "FLEX"
            filas.append(fila)
    pd.DataFrame(filas).to_csv(SALIDA_FLEX, index=False, encoding="utf-8-sig")
    return flex_por_perfil


def exportar_excel(equipos, flex_por_perfil):
    filas = []
    for perfil in ["SEGURO", "INTERMEDIO", "ARRIESGADO"]:
        for tipo, jugadores in [
            ("TITULAR", equipos.get(perfil, [])),
            ("FLEX", flex_por_perfil.get(perfil, [])),
        ]:
            for j in jugadores:
                filas.append({
                    "Perfil": perfil,
                    "Tipo": tipo,
                    "Posición": j.get("position", ""),
                    "Jugador": j.get("player_name", ""),
                    "Club": j.get("team_name", ""),
                    "Rival": j.get("rival", ""),
                    "Local": j.get("es_local", ""),
                    "Matchup score": j.get("matchup_score", ""),
                    "Contexto": motor.generar_contexto_jugador(j),
                    "Predicción Modelo C": j.get("prediccion_modelo_c", ""),
                    "Historial": j.get("partidos_historicos", ""),
                    "Participaciones últimos 3": j.get("participaciones_ultimos_3", ""),
                    "Titulares últimos 3": j.get("titulares_ultimos_3", ""),
                    "Promedio": j.get("promedio", ""),
                    "P90": j.get("p90", ""),
                    "Factor confianza": j.get("factor_confianza", ""),
                    "Sim P50": j.get("pre_sim_p50_ajustada", ""),
                    "Sim P75": j.get("pre_sim_p75_ajustada", ""),
                    "Sim P90": j.get("pre_sim_p90_ajustada", ""),
                    "Factor matchup": j.get("pre_sim_factor_matchup", ""),
                })

    df = pd.DataFrame(filas)
    columnas_numericas = [
        "Matchup score", "Predicción Modelo C", "Promedio", "P90",
        "Factor confianza", "Sim P50", "Sim P75", "Sim P90", "Factor matchup"
    ]
    for col in columnas_numericas:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce").round(3)

    with pd.ExcelWriter(SALIDA_EXCEL, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Equipos")
        ws = writer.book["Equipos"]

        # Separación visual entre SEGURO, INTERMEDIO y ARRIESGADO.
        filas_separacion = []
        for fila in range(ws.max_row, 2, -1):
            if ws.cell(fila, 1).value != ws.cell(fila - 1, 1).value:
                filas_separacion.append(fila)
        for fila in filas_separacion:
            ws.insert_rows(fila, 1)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for fila in range(2, ws.max_row + 1):
            ws.row_dimensions[fila].height = 65

        # Anchos personalizados recuperados del Excel anterior.
        anchos = {
            "Jugador": 22.25,
            "Club": 24,
            "Rival": 24,
            "Contexto": 90,
        }
        for nombre_columna, ancho in anchos.items():
            for celda in ws[1]:
                if celda.value == nombre_columna:
                    ws.column_dimensions[celda.column_letter].width = ancho
                    break

        borde_fino = Side(style="thin")
        borde_grueso = Side(style="medium")
        rellenos = {
            "SEGURO": PatternFill(fill_type="solid", fgColor="E2F0D9"),
            "INTERMEDIO": PatternFill(fill_type="solid", fgColor="FFF2CC"),
            "ARRIESGADO": PatternFill(fill_type="solid", fgColor="F4CCCC"),
        }

        # Fondo + borde fino en cada celda y borde grueso alrededor de cada perfil.
        max_col = ws.max_column
        fila_inicio = None
        perfil_actual = None

        for fila in range(2, ws.max_row + 2):
            perfil = ws.cell(fila, 1).value if fila <= ws.max_row else None

            if perfil_actual is None and perfil in rellenos:
                fila_inicio = fila
                perfil_actual = perfil

            if perfil_actual is not None and perfil != perfil_actual:
                for f in range(fila_inicio, fila):
                    for col in range(1, max_col + 1):
                        celda = ws.cell(f, col)
                        celda.fill = rellenos[perfil_actual]
                        celda.border = Border(
                            left=borde_grueso if col == 1 else borde_fino,
                            right=borde_grueso if col == max_col else borde_fino,
                            top=borde_grueso if f == fila_inicio else borde_fino,
                            bottom=borde_grueso if f == fila - 1 else borde_fino,
                        )

                fila_inicio = None
                perfil_actual = None

                if perfil in rellenos:
                    fila_inicio = fila
                    perfil_actual = perfil

        for fila in range(2, ws.max_row + 1):
            if ws.cell(fila, 1).value is None:
                continue
            ws.row_dimensions[fila].height = 65
            for col in range(1, ws.max_column + 1):
                ws.cell(fila, col).alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )

        for celda in ws[1]:
            celda.font = Font(bold=True)
            celda.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
        ws.row_dimensions[1].height = 30


def main():
    print("=" * 72)
    print(f"WINNING AI - PRE-SIMULACION + OPTIMIZADOR - FECHA {FECHA_OBJETIVO}")
    print("=" * 72)
    print("Corte historico:", CORTE.strftime("%d/%m/%Y"))
    print("Simulaciones por jugador:", N_SIM)

    historico = pd.read_csv(motor.HISTORICO_FILE, low_memory=False)
    historico["date"] = pd.to_datetime(historico["date"], errors="coerce")
    partidos = motor.cargar_partidos_fecha10()
    jugadores = motor.construir_jugadores_objetivo(historico, partidos)
    mapa = motor.construir_mapa_lineups_historicos()
    contexto, forma, local_visitante, rendimiento = motor.cargar_contextos()
    posiciones = motor.cargar_posiciones()

    candidatos = motor.construir_candidatos(
        historico, jugadores, mapa, contexto, forma,
        local_visitante, rendimiento, posiciones
    )

    modelo, columnas = motor.entrenar_modelo_c()
    candidatos = motor.agregar_prediccion_modelo_c(candidatos, modelo, columnas)
    candidatos = agregar_pre_simulacion(candidatos, historico)
    candidatos.to_csv(SALIDA_CANDIDATOS, index=False, encoding="utf-8-sig")

    activar_score_pre_simulacion()
    equipos = motor.optimizar_tres_equipos_globalmente(candidatos.copy())
    guardar_equipos(equipos)
    flex = construir_flex(candidatos, equipos)
    exportar_excel(equipos, flex)

    print()
    print("=" * 72)
    print(f"PREDICCION PRE-SIMULACION FECHA {FECHA_OBJETIVO}")
    print("=" * 72)

    for perfil in ["SEGURO", "INTERMEDIO", "ARRIESGADO"]:
        print()
        print(">>>", perfil)
        for j in equipos.get(perfil, []):
            print(
                j["position"], "|", j["player_name"], "|", j["team_name"],
                "| rival:", j.get("rival", ""),
                "| objetivo:", round(float(j.get("score_seleccion", 0)), 2)
            )
        print("FLEX:")
        for j in flex.get(perfil, []):
            print(
                j["position"], "|", j["player_name"], "|", j["team_name"],
                "| score FLEX:", round(float(j.get("score_flex", 0)), 2)
            )

    print()
    print("Archivos generados:")
    print("-", SALIDA_CANDIDATOS)
    print("-", SALIDA_EQUIPOS)
    print("-", SALIDA_FLEX)
    print("-", SALIDA_EXCEL)
    print(f"Los puntos reales de Fecha {FECHA_OBJETIVO} NO se usan para seleccionar.")


if __name__ == "__main__":
    main()
