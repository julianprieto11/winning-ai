import argparse
import json
import math
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


ROOT = Path(__file__).resolve().parent
DATOS = ROOT / "datos"
PARTIDOS_DIR = DATOS / "partidos"
DATASET = DATOS / "dataset_winning_pitchapi.csv"

SOFA_SEASON_ID = 87913
CLAUSURA_INICIO = pd.Timestamp("2026-07-23")
MAX_GOLES = 7


def cargar_partidos():
    partidos = []
    if not PARTIDOS_DIR.exists():
        return partidos

    for ruta in PARTIDOS_DIR.glob("*.json"):
        try:
            payload = json.loads(ruta.read_text(encoding="utf-8"))
            evento = payload.get("event", {})
            if not isinstance(evento, dict):
                continue
            if "event" in evento and isinstance(evento["event"], dict):
                evento = evento["event"]

            season = evento.get("season", {}) or {}
            ts = evento.get("startTimestamp")
            if not ts:
                continue

            fecha = pd.Timestamp.fromtimestamp(int(ts), tz="UTC").tz_localize(None).normalize()
            if fecha < CLAUSURA_INICIO:
                continue
            if season.get("id") not in (None, SOFA_SEASON_ID):
                continue

            round_info = evento.get("roundInfo", {}) or {}
            ronda = round_info.get("round")
            if ronda is None:
                continue

            home = evento.get("homeTeam", {}) or {}
            away = evento.get("awayTeam", {}) or {}
            score = evento.get("homeScore", {}) or {}
            away_score = evento.get("awayScore", {}) or {}

            home_goals = score.get("current")
            away_goals = away_score.get("current")

            partidos.append({
                "match_id": str(evento.get("id") or ruta.stem),
                "fecha": fecha,
                "ronda": int(ronda),
                "local": home.get("name", ""),
                "visitante": away.get("name", ""),
                "goles_local": float(home_goals) if home_goals is not None else None,
                "goles_visitante": float(away_goals) if away_goals is not None else None,
            })
        except Exception:
            continue

    return pd.DataFrame(partidos).drop_duplicates("match_id") if partidos else pd.DataFrame()


def normalizar_equipo(nombre):
    texto = str(nombre or "").strip().lower()
    reemplazos = {
        "instituto": "instituto de cordoba",
        "instituto de cordoba": "instituto de cordoba",
        "independiente": "ca independiente",
        "ca independiente": "ca independiente",
        "union": "club atletico union de santa fe",
        "club atletico union de santa fe": "club atletico union de santa fe",
        "talleres": "ca talleres",
        "ca talleres": "ca talleres",
        "estudiantes": "estudiantes de la plata",
        "estudiantes de la plata": "estudiantes de la plata",
        "belgrano": "club atletico belgrano",
        "club atletico belgrano": "club atletico belgrano",
        "lanus": "ca lanus",
        "ca lanus": "ca lanus",
        "gimnasia lp": "gimnasia y esgrima",
        "gimnasia y esgrima": "gimnasia y esgrima",
        "central cordoba de santiago": "central cordoba",
        "central cordoba": "central cordoba",
    }
    import unicodedata
    texto = unicodedata.normalize("NFD", texto)
    texto = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    return reemplazos.get(texto, texto)


def poisson(k, lamb):
    if lamb <= 0:
        return 1.0 if k == 0 else 0.0
    return math.exp(-lamb) * (lamb ** k) / math.factorial(k)


def construir_fuerzas(historicos):
    rows = []
    for equipo in sorted(set(historicos["local"]).union(historicos["visitante"])):
        h = historicos[historicos["local"] == equipo]
        a = historicos[historicos["visitante"] == equipo]
        todos = pd.concat([
            h.assign(gf=h["goles_local"], ga=h["goles_visitante"]),
            a.assign(gf=a["goles_visitante"], ga=a["goles_local"]),
        ], ignore_index=True)
        todos = todos.dropna(subset=["gf", "ga"])

        recent_ids = todos.sort_values("fecha").tail(5).index
        recent = todos.loc[recent_ids]

        rows.append({
            "equipo": equipo,
            "pj": len(todos),
            "gf_pj": todos["gf"].mean() if len(todos) else 0.0,
            "ga_pj": todos["ga"].mean() if len(todos) else 0.0,
            "gf_reciente": recent["gf"].mean() if len(recent) else 0.0,
            "ga_reciente": recent["ga"].mean() if len(recent) else 0.0,
            "local_gf": h["goles_local"].mean() if len(h) else float("nan"),
            "local_ga": h["goles_visitante"].mean() if len(h) else float("nan"),
            "visit_gf": a["goles_visitante"].mean() if len(a) else float("nan"),
            "visit_ga": a["goles_local"].mean() if len(a) else float("nan"),
        })
    return pd.DataFrame(rows)


def predecir_partido(row, fuerzas, historicos):
    liga = historicos.dropna(subset=["goles_local", "goles_visitante"])
    home_avg = liga["goles_local"].mean() or 1.0
    away_avg = liga["goles_visitante"].mean() or 1.0

    def get(team):
        x = fuerzas[fuerzas["equipo_norm"] == normalizar_equipo(team)]
        return x.iloc[0] if not x.empty else None

    h = get(row["local"])
    a = get(row["visitante"])

    if h is None or a is None:
        return None

    h_attack = ((h["local_gf"] if pd.notna(h["local_gf"]) else h["gf_pj"]) / home_avg)
    h_def = ((h["local_ga"] if pd.notna(h["local_ga"]) else h["ga_pj"]) / away_avg)
    a_attack = ((a["visit_gf"] if pd.notna(a["visit_gf"]) else a["gf_pj"]) / away_avg)
    a_def = ((a["visit_ga"] if pd.notna(a["visit_ga"]) else a["ga_pj"]) / home_avg)

    xg_home = home_avg * h_attack * a_def
    xg_away = away_avg * a_attack * h_def

    # 30% de forma reciente, 70% de rendimiento estructural.
    xg_home *= 0.70 + 0.30 * max(0.35, min(1.65, ((h["gf_reciente"] + a["ga_reciente"]) / 2) / max(0.01, (h["gf_pj"] + a["ga_pj"]) / 2)))
    xg_away *= 0.70 + 0.30 * max(0.35, min(1.65, ((a["gf_reciente"] + h["ga_reciente"]) / 2) / max(0.01, (a["gf_pj"] + h["ga_pj"]) / 2)))

    xg_home = max(0.05, min(4.5, xg_home))
    xg_away = max(0.05, min(4.5, xg_away))

    matriz = []
    for gh in range(MAX_GOLES + 1):
        for ga in range(MAX_GOLES + 1):
            p = poisson(gh, xg_home) * poisson(ga, xg_away)
            matriz.append((gh, ga, p))

    total = sum(p for _, _, p in matriz) or 1.0
    prob_local = sum(p for gh, ga, p in matriz if gh > ga) / total
    prob_empate = sum(p for gh, ga, p in matriz if gh == ga) / total
    prob_visitante = sum(p for gh, ga, p in matriz if gh < ga) / total
    mejor = max(matriz, key=lambda x: x[2])

    return {
        "local": row["local"],
        "visitante": row["visitante"],
        "xG_local": xg_home,
        "xG_visitante": xg_away,
        "prob_local": prob_local,
        "prob_empate": prob_empate,
        "prob_visitante": prob_visitante,
        "marcador_mas_probable": f"{mejor[0]}-{mejor[1]}",
        "prob_marcador": mejor[2],
        "resultado_mas_probable": max(
            [("LOCAL", prob_local), ("EMPATE", prob_empate), ("VISITANTE", prob_visitante)],
            key=lambda x: x[1],
        )[0],
    }


def guardar_excel(df, ruta, fecha, corte):
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Predicciones"
    ws.freeze_panes = "A4"

    ws.merge_cells("A1:K1")
    ws["A1"] = f"WINNING AI — PREDICCIÓN DE RESULTADOS FECHA {fecha}"
    ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="001E5F")
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:K2")
    ws["A2"] = f"Modelo Poisson contextual | Datos históricos hasta {corte}"
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = ["Local", "Visitante", "xG Local", "xG Visitante", "Prob. Local", "Prob. Empate", "Prob. Visitante", "Resultado más probable", "Marcador más probable", "Prob. marcador", "PJ histórico"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(4, col, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="001E5F")
        c.alignment = Alignment(horizontal="center", wrap_text=True)

    for r, (_, row) in enumerate(df.iterrows(), 5):
        vals = [
            row["local"], row["visitante"], row["xG_local"], row["xG_visitante"],
            row["prob_local"], row["prob_empate"], row["prob_visitante"],
            row["resultado_mas_probable"], row["marcador_mas_probable"],
            row["prob_marcador"], row["PJ histórico"],
        ]
        for c, v in enumerate(vals, 1):
            ws.cell(r, c, v)
        for c in (3, 4, 10):
            ws.cell(r, c).number_format = "0.00"
        for c in (5, 6, 7):
            ws.cell(r, c).number_format = "0.0%"

    widths = [26, 26, 12, 14, 14, 14, 16, 22, 22, 14, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64+i)].width = w

    wb.save(ruta)


def main(fecha):
    partidos = cargar_partidos()
    if partidos.empty:
        raise RuntimeError("No hay archivos válidos en datos/partidos. Ejecutá primero ejecutar_fecha.py N.")

    objetivo = partidos[partidos["ronda"] == int(fecha)].copy()
    if objetivo.empty:
        raise RuntimeError(f"No se encontraron partidos de la Fecha {fecha} en datos/partidos.")

    inicio = objetivo["fecha"].min()
    historicos = partidos[
        (partidos["fecha"] < inicio)
        & partidos["goles_local"].notna()
        & partidos["goles_visitante"].notna()
    ].copy()

    if historicos.empty:
        raise RuntimeError("No hay suficientes partidos históricos para construir probabilidades.")

    fuerzas = construir_fuerzas(historicos)
    fuerzas["equipo_norm"] = fuerzas["equipo"].map(normalizar_equipo)

    resultados = []
    for _, row in objetivo.sort_values(["fecha", "local"]).iterrows():
        pred = predecir_partido(row, fuerzas, historicos)
        if pred is None:
            continue
        pred["PJ histórico"] = int(
            fuerzas.loc[
                fuerzas["equipo_norm"].isin([
                    normalizar_equipo(row["local"]),
                    normalizar_equipo(row["visitante"]),
                ]),
                "pj",
            ].sum()
        )
        resultados.append(pred)

    if not resultados:
        raise RuntimeError("No se pudo predecir ningún partido.")

    salida = pd.DataFrame(resultados)
    salida["fecha"] = int(fecha)
    salida["corte_historico"] = inicio.strftime("%Y-%m-%d")

    out_dir = DATOS / "predicciones" / f"fecha_{int(fecha)}"
    out_dir.mkdir(parents=True, exist_ok=True)
    csv = out_dir / "predicciones_resultados.csv"
    xlsx = out_dir / "predicciones_resultados.xlsx"
    salida.to_csv(csv, index=False, encoding="utf-8-sig")
    guardar_excel(salida, xlsx, fecha, inicio.strftime("%Y-%m-%d"))

    print("=" * 90)
    print(f"PREDICCIÓN DE RESULTADOS — FECHA {fecha}")
    print("=" * 90)
    print(f"Partidos objetivo: {len(objetivo)}")
    print(f"Partidos históricos usados: {len(historicos)}")
    print(f"Corte histórico: {inicio.date()}")
    print(f"CSV: {csv}")
    print(f"Excel: {xlsx}")
    print()
    print(salida[["local", "visitante", "prob_local", "prob_empate", "prob_visitante", "resultado_mas_probable", "marcador_mas_probable"]].to_string(index=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Predicción probabilística de resultados por fecha.")
    parser.add_argument("fecha", type=int, help="Fecha a predecir. Ej.: 12")
    args = parser.parse_args()
    main(args.fecha)
