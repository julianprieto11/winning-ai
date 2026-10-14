"""Analiza duelos desparejos entre defensores y atacantes para una fecha.
Uso: python analizar_duelos_desparejos.py --fecha 13
No modifica el motor de equipos, el aprendizaje ni el predictor de partidos.
Los cruces son estimaciones tácticas: no representan marcajes individuales confirmados.
"""
from __future__ import annotations

import argparse
import json
import re
import unicodedata
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
DATOS = BASE / "datos"
PARTIDOS_DIR = DATOS / "partidos"
HISTORICO = DATOS / "dataset_winning_pitchapi.csv"
LINEUPS_DIR = DATOS / "pitchapi" / "lineups"
SALIDA = DATOS / "duelos_desparejos_fecha{fecha}.xlsx"
INICIO_CLAUSURA = pd.Timestamp("2026-07-23")

# Métricas por 90 minutos. Los pesos comparan fortalezas relevantes para cada lado.
METRICAS_DEF = {
    "duels_won": ("Duelos ganados", 1.0),
    "interceptions": ("Intercepciones", 1.0),
    "tackles": ("Entradas / tackles", 0.8),
    "recoveries": ("Recuperaciones", 0.6),
    "clearances": ("Despejes", 0.5),
    "blocks": ("Bloqueos", 0.5),
}
METRICAS_DEL = {
    "take_ons_won": ("Regates exitosos", 1.0),
    "duels_won": ("Duelos ganados", 0.9),
    "progressive_carries": ("Conducciones progresivas", 0.8),
    "fouls_won": ("Faltas recibidas", 0.6),
    "chances_created": ("Ocasiones creadas", 0.6),
    "shots_on_target": ("Tiros al arco", 0.5),
}
MIN_MINUTOS_MUESTRA = 45
ULTIMAS_APARICIONES = 5


def normalizar(valor):
    texto = unicodedata.normalize("NFKD", str(valor or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", texto.lower()).strip()


def encontrar_evento(obj):
    if isinstance(obj, dict):
        e = obj.get("event")
        if isinstance(e, dict) and "homeTeam" in e and "awayTeam" in e:
            return e
        for v in obj.values():
            encontrado = encontrar_evento(v)
            if encontrado:
                return encontrado
    elif isinstance(obj, list):
        for v in obj:
            encontrado = encontrar_evento(v)
            if encontrado:
                return encontrado
    return None


def partidos_objetivo(fecha):
    filas = []
    if not PARTIDOS_DIR.exists():
        raise FileNotFoundError(f"No existe la carpeta {PARTIDOS_DIR}")
    for path in PARTIDOS_DIR.glob("*.json"):
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            continue
        e = encontrar_evento(payload)
        if not e:
            continue
        round_info = e.get("roundInfo") or {}
        round_no = round_info.get("round")
        ts = e.get("startTimestamp")
        fecha_partido = pd.to_datetime(ts, unit="s", errors="coerce") if ts else pd.NaT
        if pd.isna(fecha_partido) or fecha_partido < INICIO_CLAUSURA or str(round_no) != str(fecha):
            continue
        home = (e.get("homeTeam") or {}).get("name")
        away = (e.get("awayTeam") or {}).get("name")
        if not home or not away:
            continue
        filas.append({
            "match_id": str(e.get("id") or path.stem),
            "fecha_partido": fecha_partido,
            "local": str(home),
            "visitante": str(away),
        })
    if not filas:
        raise ValueError(
            f"No encontré partidos de la Fecha {fecha} del Clausura en datos/partidos/*.json. "
            "Verificá que los JSON estén actualizados y que roundInfo.round corresponda a esa fecha."
        )
    return pd.DataFrame(filas).drop_duplicates("match_id").sort_values("fecha_partido")


def cargar_lineup(match_id):
    path = LINEUPS_DIR / f"{match_id}_lineups.json"
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return None
    data = payload.get("data")
    if not isinstance(data, dict):
        return None
    if data.get("confirmed") is False:
        home_ok = bool((data.get("home") or {}).get("confirmed"))
        away_ok = bool((data.get("away") or {}).get("confirmed"))
        if not (home_ok and away_ok):
            return None
    return data


def nombre_lado(lineup, club):
    objetivo = normalizar(club)
    for lado in ("home", "away"):
        equipo = (lineup.get(f"{lado}_team") or {}).get("name", "")
        if normalizar(equipo) == objetivo:
            return lado
    # No se adivina un equipo si el nombre no coincide.
    return None


def player_id(item):
    value = item.get("player_id", item.get("id"))
    return str(value) if value is not None else ""


def role_text(item):
    partes = []
    for key in ("position_name", "positionName", "position", "role", "formation_position", "formationPosition", "position_label"):
        value = item.get(key)
        if isinstance(value, str):
            partes.append(value)
        elif isinstance(value, dict):
            partes.extend(str(value.get(k, "")) for k in ("name", "label", "short_name", "shortName", "abbreviation"))
    return normalizar(" ".join(partes))


def sector_rol(texto, grupo):
    """Solo asigna sector cuando el JSON aporta una posición textual explícita."""
    t = f" {normalizar(texto)} "
    derecha = any(x in t for x in (" right ", " derecha ", " rb ", " rwb ", " right back ", " right winger ", " rw "))
    izquierda = any(x in t for x in (" left ", " izquierda ", " lb ", " lwb ", " left back ", " left winger ", " lw "))
    if grupo == "DEF":
        if any(x in t for x in (" center back ", " centre back ", " central defender ", " cb ", " defensa central ")):
            return "CENTRAL"
        if derecha and any(x in t for x in ("back", "wing back", "lateral", " rb ", " rwb ")):
            return "DERECHA"
        if izquierda and any(x in t for x in ("back", "wing back", "lateral", " lb ", " lwb ")):
            return "IZQUIERDA"
    if grupo == "DEL":
        if any(x in t for x in ("striker", "centre forward", "center forward", "central forward", " delantero centro ", " cf ", " st ")):
            return "CENTRAL"
        if derecha and any(x in t for x in ("wing", "rw", "right forward")):
            return "DERECHA"
        if izquierda and any(x in t for x in ("wing", "lw", "left forward")):
            return "IZQUIERDA"
    return None


def lineup_players(lineup, club):
    if not lineup:
        return None
    lado = nombre_lado(lineup, club)
    if not lado:
        return None
    side_data = lineup.get(lado) or {}
    starters = side_data.get("starters")
    if not isinstance(starters, list) or not starters:
        return None
    result = {}
    for item in starters:
        if not isinstance(item, dict):
            continue
        pid = player_id(item)
        if pid:
            result[pid] = {
                "role_text": role_text(item),
                "name": str(item.get("player_name") or item.get("name") or item.get("short_name") or ""),
            }
    return result


def preparar_historico(df):
    alias = {
        "player_id": ["player_id"],
        "nombre": ["player_name", "nombre_jugador", "name"],
        "equipo": ["team_name", "club", "team", "equipo"],
        "partido": ["match_id", "fixture_id", "partido_id"],
        "fecha": ["match_date", "date", "fecha", "fecha_partido"],
        "minutos": ["minutes_played", "minutes"],
    }
    cols = {}
    for destino, opciones in alias.items():
        cols[destino] = next((c for c in opciones if c in df.columns), None)
    faltan = [k for k, v in cols.items() if v is None]
    if faltan:
        raise ValueError("Al histórico le faltan columnas: " + ", ".join(faltan))
    df = df.copy()
    for destino, original in cols.items():
        if destino != original:
            df[destino] = df[original]
    df["player_id"] = df["player_id"].astype(str)
    df["partido"] = df["partido"].astype(str)
    df["equipo"] = df["equipo"].astype(str)
    df["fecha"] = pd.to_datetime(df["fecha"], errors="coerce", utc=True).dt.tz_convert(None)
    df["minutos"] = pd.to_numeric(df["minutos"], errors="coerce").fillna(0)
    for col in set(METRICAS_DEF) | set(METRICAS_DEL):
        if col not in df:
            df[col] = np.nan
        df[col] = pd.to_numeric(df[col], errors="coerce")
    return df


def lista_partidos_equipo(df, equipo):
    return (
        df[df["equipo"].map(normalizar) == normalizar(equipo)][["partido", "fecha"]]
        .dropna().drop_duplicates().sort_values(["fecha", "partido"], ascending=False)
    )


def validar_titularidad_2de2(df, player, team, cache):
    """Devuelve elegibilidad, rol reciente y motivo. Exige lineup confirmado en ambos partidos."""
    partidos = lista_partidos_equipo(df, team).head(2)
    if len(partidos) < 2:
        return False, None, "No hay dos partidos previos verificables del equipo"
    roles = []
    for _, row in partidos.iterrows():
        mid = str(row["partido"])
        if mid not in cache:
            cache[mid] = cargar_lineup(mid)
        lineup = cache[mid]
        if lineup is None:
            return False, None, f"Falta alineación confirmada de uno de los últimos 2 partidos ({mid})"
        players = lineup_players(lineup, team)
        if players is None:
            return False, None, f"No se pudo verificar el equipo en la alineación ({mid})"
        if str(player) not in players:
            return False, None, "No fue titular en los últimos 2 partidos (exigencia 2/2)"
        roles.append(players[str(player)]["role_text"])
    # Se usa el rol más reciente; si los roles textuales discrepan, no se inventa uno.
    if roles[0] != roles[1]:
        return True, None, "Titular 2/2, pero cambió/no se pudo confirmar el mismo sector táctico"
    return True, roles[0], "Titular confirmado 2/2"


def metricas_por90(df, pid, team):
    bloque = df[(df["player_id"] == str(pid)) & (df["equipo"].map(normalizar) == normalizar(team))]
    bloque = bloque.sort_values(["fecha", "partido"]).drop_duplicates("partido", keep="last").tail(ULTIMAS_APARICIONES)
    minutos = float(bloque["minutos"].sum())
    if minutos < MIN_MINUTOS_MUESTRA:
        return None
    resultado = {"minutos_muestra": minutos, "partidos_muestra": int(bloque["partido"].nunique())}
    for col in set(METRICAS_DEF) | set(METRICAS_DEL):
        total = bloque[col].sum(min_count=1)
        resultado[col] = float(total) / minutos * 90 if pd.notna(total) else np.nan
    nombres = bloque["nombre"].dropna().astype(str)
    resultado["nombre"] = nombres.iloc[-1] if len(nombres) else str(pid)
    return resultado


def normalizar_vector(datos, metricas):
    vals = []
    for col, (_, peso) in metricas.items():
        val = datos.get(col, np.nan)
        if pd.notna(val):
            vals.append((max(0.0, float(val)), peso))
    if not vals:
        return np.nan
    # Escala logarítmica: evita que una estadística de mucho volumen domine todo.
    return sum(np.log1p(v) * w for v, w in vals) / sum(w for _, w in vals)


def detalle_metricas(datos, metricas):
    return "; ".join(
        f"{nombre}: {datos[col]:.2f}/90"
        for col, (nombre, _) in metricas.items()
        if col in datos and pd.notna(datos[col])
    )


def compatibilidad(sector_def, sector_del):
    # Cruces laterales cruzados: extremo izquierdo vs lateral derecho y viceversa.
    if sector_def == "CENTRAL" and sector_del == "CENTRAL":
        return "CENTRAL vs CENTRAL"
    if sector_def == "DERECHA" and sector_del == "IZQUIERDA":
        return "Extremo izquierdo vs lateral derecho"
    if sector_def == "IZQUIERDA" and sector_del == "DERECHA":
        return "Extremo derecho vs lateral izquierdo"
    return None


def main():
    parser = argparse.ArgumentParser(description="Detecta duelos estimados desparejos DEF vs DEL.")
    parser.add_argument("--fecha", type=int, required=True, help="Número de fecha del Clausura")
    args = parser.parse_args()

    if not HISTORICO.exists():
        raise FileNotFoundError(f"No existe el histórico: {HISTORICO}")
    print(f"Cargando histórico PitchAPI: {HISTORICO}")
    historico = preparar_historico(pd.read_csv(HISTORICO, low_memory=False))
    fixtures = partidos_objetivo(args.fecha)
    cache_lineups = {}
    salidas = []
    descartes = []

    # Se toman jugadores que figuran en el histórico de cada club; no se inventan planteles.
    for _, fixture in fixtures.iterrows():
        home, away = fixture["local"], fixture["visitante"]
        for equipo_ataca, equipo_defiende in ((home, away), (away, home)):
            atacantes = historico[historico["equipo"].map(normalizar) == normalizar(equipo_ataca)]["player_id"].dropna().astype(str).unique()
            defensores = historico[historico["equipo"].map(normalizar) == normalizar(equipo_defiende)]["player_id"].dropna().astype(str).unique()
            elegibles_del = []
            elegibles_def = []
            for pid in atacantes:
                ok, role, motivo = validar_titularidad_2de2(historico, pid, equipo_ataca, cache_lineups)
                stats = metricas_por90(historico, pid, equipo_ataca) if ok else None
                if ok and stats and role is not None:
                    sector = sector_rol(role, "DEL")
                    if sector:
                        elegibles_del.append((pid, stats, sector, role))
                    else:
                        descartes.append({"Partido": f"{home} vs {away}", "Jugador": stats["nombre"], "Equipo": equipo_ataca, "Motivo": "Rol de delantero sin sector explícito en lineup"})
                elif ok and stats:
                    descartes.append({"Partido": f"{home} vs {away}", "Jugador": stats["nombre"], "Equipo": equipo_ataca, "Motivo": motivo})
            for pid in defensores:
                ok, role, motivo = validar_titularidad_2de2(historico, pid, equipo_defiende, cache_lineups)
                stats = metricas_por90(historico, pid, equipo_defiende) if ok else None
                if ok and stats and role is not None:
                    sector = sector_rol(role, "DEF")
                    if sector:
                        elegibles_def.append((pid, stats, sector, role))
                    else:
                        descartes.append({"Partido": f"{home} vs {away}", "Jugador": stats["nombre"], "Equipo": equipo_defiende, "Motivo": "Rol de defensor sin sector explícito en lineup"})
                elif ok and stats:
                    descartes.append({"Partido": f"{home} vs {away}", "Jugador": stats["nombre"], "Equipo": equipo_defiende, "Motivo": motivo})

            for pid_del, del_stats, sector_del, rol_del in elegibles_del:
                for pid_def, def_stats, sector_def, rol_def in elegibles_def:
                    cruce = compatibilidad(sector_def, sector_del)
                    if not cruce:
                        continue
                    fuerza_def = normalizar_vector(def_stats, METRICAS_DEF)
                    fuerza_del = normalizar_vector(del_stats, METRICAS_DEL)
                    if pd.isna(fuerza_def) or pd.isna(fuerza_del):
                        continue
                    ventaja = fuerza_def - fuerza_del
                    if abs(ventaja) < 0.12:
                        continue
                    if ventaja > 0:
                        dominador = def_stats["nombre"]
                        tipo = "Defensor con ventaja estadística"
                    else:
                        dominador = del_stats["nombre"]
                        tipo = "Atacante con ventaja estadística"
                    salidas.append({
                        "Fecha": args.fecha,
                        "Partido": f"{home} vs {away}",
                        "Fecha partido": fixture["fecha_partido"].strftime("%d/%m/%Y"),
                        "Equipo defensor": equipo_defiende,
                        "Defensor": def_stats["nombre"],
                        "Rol defensor": rol_def,
                        "Equipo atacante": equipo_ataca,
                        "Atacante": del_stats["nombre"],
                        "Rol atacante": rol_del,
                        "Sector compatible": cruce,
                        "Ventaja": tipo,
                        "Jugador favorecido": dominador,
                        "Índice defensor": round(fuerza_def, 3),
                        "Índice atacante": round(fuerza_del, 3),
                        "Diferencia índice": round(ventaja, 3),
                        "Muestra defensor": f'{def_stats["partidos_muestra"]} partidos / {def_stats["minutos_muestra"]:.0f} min',
                        "Muestra atacante": f'{del_stats["partidos_muestra"]} partidos / {del_stats["minutos_muestra"]:.0f} min',
                        "Stats defensor por 90": detalle_metricas(def_stats, METRICAS_DEF),
                        "Stats atacante por 90": detalle_metricas(del_stats, METRICAS_DEL),
                        "Advertencia": "DUELO ESTIMADO por sector; no confirma marcaje individual",
                    })

    columnas = [
        "Fecha", "Partido", "Fecha partido", "Equipo defensor", "Defensor", "Rol defensor",
        "Equipo atacante", "Atacante", "Rol atacante", "Sector compatible", "Ventaja",
        "Jugador favorecido", "Índice defensor", "Índice atacante", "Diferencia índice",
        "Muestra defensor", "Muestra atacante", "Stats defensor por 90",
        "Stats atacante por 90", "Advertencia",
    ]
    resultado = pd.DataFrame(salidas, columns=columnas)
    if not resultado.empty:
        resultado["_abs"] = resultado["Diferencia índice"].abs()
        resultado = resultado.sort_values(["_abs", "Partido"], ascending=[False, True]).drop(columns="_abs")

    resumen = pd.DataFrame([{
        "Fecha analizada": args.fecha,
        "Partidos encontrados": len(fixtures),
        "Duelos compatibles encontrados": len(resultado),
        "Regla titularidad": "Ambos jugadores titulares confirmados en los últimos 2 de 2 partidos de su equipo",
        "Regla sector": "Solo cruces centrales o laterales opuestos con rol textual explícito",
        "Ventana estadística": f"Últimas {ULTIMAS_APARICIONES} apariciones; mínimo {MIN_MINUTOS_MUESTRA} minutos",
        "Nota": "Si no hay sector táctico verificable, el jugador no entra en el ranking.",
    }])
    descartes_df = pd.DataFrame(descartes, columns=["Partido", "Jugador", "Equipo", "Motivo"])

    salida = SALIDA.with_name(SALIDA.name.format(fecha=args.fecha))
    salida.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(salida, engine="openpyxl") as writer:
        resumen.to_excel(writer, sheet_name="Resumen", index=False)
        resultado.to_excel(writer, sheet_name="Duelos", index=False)
        descartes_df.to_excel(writer, sheet_name="Exclusiones", index=False)

    wb = load_workbook(salida)
    navy = "001E5F"
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 32
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            width = min(max(max(len(str(c.value or "")) for c in col[:100]) + 2, 12), 48)
            ws.column_dimensions[letter].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(salida)

    print(f"Partidos de la Fecha {args.fecha}: {len(fixtures)}")
    print(f"Duelos que cumplen TODOS los filtros: {len(resultado)}")
    print(f"Excel generado: {salida}")
    if resultado.empty:
        print("No se relajaron filtros. Revisá la hoja 'Exclusiones' y los roles textuales de las alineaciones.")
    else:
        print(resultado[["Partido", "Defensor", "Atacante", "Sector compatible", "Ventaja"]].head(10).to_string(index=False))


if __name__ == "__main__":
    main()
